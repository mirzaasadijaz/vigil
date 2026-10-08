"""Read .md / .txt / .pdf / .docx from data/rag/documents/<category>/..."""
from __future__ import annotations

import hashlib
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator

from backend.rag import DOCS_DIR

SUPPORTED = {".md", ".txt", ".pdf", ".docx"}


@dataclass
class Document:
    source: str        # path relative to DOCS_DIR, forward slashes (stable id)
    category: str      # first sub-folder: sop / policies / manuals / general
    text: str
    sha1: str          # of the file bytes, used to skip unchanged files


def _read(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in (".md", ".txt"):
        return path.read_text(encoding="utf-8", errors="replace")
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
    if ext == ".docx":
        import docx
        return "\n\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    return ""


def iter_documents(root: Path | None = None) -> Iterator[Document]:
    root = Path(root or DOCS_DIR)
    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED or path.name.startswith("."):
            continue
        rel = path.relative_to(root)
        try:
            text = _read(path).strip()
        except Exception as e:                       # one bad PDF must not stop the whole ingest
            print(f"[RAG] skipped {rel}: {e}", flush=True)
            continue
        if not text:
            continue
        yield Document(
            source=rel.as_posix(),
            category=rel.parts[0] if len(rel.parts) > 1 else "general",
            text=text,
            sha1=hashlib.sha1(path.read_bytes()).hexdigest(),
        )
